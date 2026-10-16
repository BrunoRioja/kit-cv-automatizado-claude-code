"""Busca vacantes en los portales de las empresas objetivo y en Bumeran Perú, y guarda en vacantes.xlsx solo las nuevas.

Uso:
    python buscar_vacantes.py                    # últimos 30 días, portales + Bumeran
    python buscar_vacantes.py --dias 7
    python buscar_vacantes.py --solo-portales    # o --sin-portales
    python buscar_vacantes.py --manual vac.json  # vacante pegada en el chat (fuera de Bumeran)
    python buscar_vacantes.py --archivar         # mueve todo Vacantes a Historial
    python buscar_vacantes.py --archivar --conservar-con-cv   # deja solo las que ya tienen CV

Toda vacante que ya esté en Vacantes o en Historial (mismo id, o misma empresa + título
aunque Bumeran la republique con otro id) se descarta como ya vista.

Prioridad al portal de la empresa: si la misma vacante sale en Bumeran, se queda la del
portal (fuente "… · también en Bumeran"). Si ya estaba en Vacantes desde Bumeran, se le
cambia el link y la fuente al portal sin tocar su calificación ni su análisis.
Los portales están en portales_empresas.py.

Usa el mismo servicio que la web pública de Bumeran y de cada portal (sin iniciar sesión).
Hace una consulta a la vez, con pausas, para uso personal y moderado.
Las columnas encaje / cubre / falta / alertas / decision las llena Claude o el usuario;
el script nunca las sobrescribe.
"""

import argparse
import re
import sys
import time
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import portales_empresas

ARCHIVO = Path(__file__).parent / "vacantes.xlsx"
API = "https://www.bumeran.com.pe/api/avisos/searchV2"
HEADERS = {
    "Content-Type": "application/json",
    "Accept": "application/json",
    "x-site-id": "BMPE",
    "Referer": "https://www.bumeran.com.pe/",
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/128.0 Safari/537.36",
}
PAUSA_SEG = 1.5
POR_PAGINA = 20
MAX_PAGINAS = 5  # por palabra clave

# ---------------------------------------------------------------------------
# Criterios de búsqueda (editar aquí)
# Ejemplo: ajústalos a los puestos que buscas (los de tus CVs son un buen punto de partida).
# ---------------------------------------------------------------------------
PALABRAS_CLAVE = {
    "Datos-BI": [
        "analista bi", "inteligencia de negocios", "business intelligence",
        "analista de datos", "data analyst", "power bi", "data bi",
        "analista de reportes", "analista de informacion",
        "automatizacion de datos", "analista de automatizacion", "power automate",
    ],
    "Comercial": [
        "analista comercial", "inteligencia comercial", "informacion comercial",
        "analista de comercializacion", "commercial analyst",
        "cuentas clave", "key account", "planeamiento comercial",
    ],
    "IA / Producto / Proyectos": [
        "inteligencia artificial", "soluciones de ia", "prompt",
        "analista de producto", "canales digitales", "analista de proyectos",
    ],
}

# Se aceptan si la ubicación contiene alguno de estos textos, o si es remoto.
UBICACIONES = ["lima", "callao", "chiclayo", "lambayeque"]

# Puestos que se descartan por título (palabras completas, sin tildes).
EXCLUIR_TITULO = [
    "jefe", "gerente", "head", "supervisor", "practicante",
    "practicas", "practica", "qa", "tester", "testing", ".net", "java",
    "vendedor", "asesor de ventas", "promotor", "call center", "teleoperador",
    "contador", "auxiliar contable", "docente", "asesor", "asesora",
    "atencion al cliente", "reclutamiento", "seleccion", "recursos humanos",
    "seguridad de la informacion", "ciberseguridad", "almacen", "soporte ti",
    "arquitecto", "arquitecto/a", "funcionario", "cajero", "gestor de caja",
    # Títulos en inglés de los portales (gerencia y prácticas; los senior sí entran)
    "manager", "director", "intern", "internship",
]

# En los portales, una vacante sin ninguna PALABRA_CLAVE en título o detalle entra igual
# si su título tiene alguna de estas (los portales titulan mucho en inglés).
TITULO_DATOS = ["data", "datos", "bi", "business intelligence", "analytics", "analitica",
                "insights", "reporting", "scientist", "cientifico de datos", "analyst",
                "analista", "product", "producto", "commercial", "comercial", "planning",
                "planeamiento", "strategy", "estrategia", "pricing"]

# Tus herramientas: si la vacante las menciona, cuentan a favor.
HERRAMIENTAS = [
    "power bi", "dax", "power query", "looker", "power automate", "power apps",
    "power platform", "appsheet", "python", "pandas", "excel", "vba",
    "apps script", "tableau", "etl", "dashboard", "kpi", "automatiz",
]

COLUMNAS = [
    # postular: "Y" = el usuario postuló/enviará (se marca aquí o en el artefacto Radar de vacantes).
    # cv_id: lo llena Claude al generar y congelar el CV; fila con cv_id = ya atendida.
    "podio", "postular", "cv_id", "id", "estado", "fuente", "fecha_publicacion", "titulo", "empresa", "ubicacion",
    "modalidad", "jornada", "carril", "palabras_clave", "coincidencias",
    "anios_pedidos", "pide_sql", "pide_ingles", "link", "encontrada_el",
    "encaje", "cubre", "falta", "alertas",
    # Resultado de la skill analizar-vacante-cv (solo para las vacantes analizadas)
    "recomendacion", "puntaje_global", "nota_empresa", "nota_puesto", "nota_encaje",
    "sueldo_esperado", "analisis_empresa", "analisis_puesto", "analisis_encaje",
    "analisis_sueldo", "veredicto", "bullets_sugeridos", "traslado", "fuentes",
    "detalle",
]
ANCHOS = {"podio": 22, "postular": 10, "cv_id": 22, "bullets_sugeridos": 40, "traslado": 22, "titulo": 45, "empresa": 30, "ubicacion": 24,
          "palabras_clave": 30, "coincidencias": 35, "link": 40, "cubre": 40, "falta": 40,
          "alertas": 40, "recomendacion": 20, "sueldo_esperado": 18, "analisis_empresa": 60,
          "analisis_puesto": 60, "analisis_encaje": 60, "analisis_sueldo": 50,
          "veredicto": 60, "fuentes": 50, "detalle": 80}


def normalizar(texto):
    texto = unicodedata.normalize("NFKD", texto or "")
    return "".join(c for c in texto if not unicodedata.combining(c)).lower()


def slug(texto):
    return re.sub(r"[^a-z0-9]+", "-", normalizar(texto)).strip("-")


def contiene_palabra(texto, palabra):
    return re.search(r"(?<![a-z0-9])" + re.escape(palabra) + r"(?![a-z0-9])", texto) is not None


def buscar(query, pagina):
    r = requests.post(
        API,
        params={"pageSize": POR_PAGINA, "page": pagina, "sort": "RECIENTES"},
        headers=HEADERS,
        json={"filtros": [], "query": query},
        timeout=30,
    )
    r.raise_for_status()
    return r.json().get("content", [])


def fecha_aviso(aviso):
    return datetime.strptime(aviso["fechaPublicacion"], "%d-%m-%Y").date()


def ubicacion_valida(ubicacion, modalidad):
    ubic, modalidad = normalizar(ubicacion), normalizar(modalidad)
    return any(r in modalidad or r in ubic for r in ("remoto", "remote")) or any(u in ubic for u in UBICACIONES)


def titulo_excluido(titulo):
    titulo = normalizar(titulo)
    return any(contiene_palabra(titulo, p) for p in EXCLUIR_TITULO)


def pasa_filtros(aviso):
    return (ubicacion_valida(aviso.get("localizacion"), aviso.get("modalidadTrabajo"))
            and not titulo_excluido(aviso.get("titulo")))


def analizar_detalle(detalle):
    texto = normalizar(detalle)
    coincidencias = [h for h in HERRAMIENTAS if h in texto]
    anios = [int(n) for n in re.findall(r"(\d{1,2})\s*(?:\+\s*)?anos?", texto) if int(n) <= 15]
    return {
        "coincidencias": ", ".join(coincidencias),
        "anios_pedidos": max(anios) if anios else None,
        "pide_sql": "sí" if contiene_palabra(texto, "sql") else "",
        "pide_ingles": "sí" if re.search(r"ingles|english", texto) else "",
    }


def recolectar(dias):
    limite = date.today() - timedelta(days=dias)
    encontrados = {}
    for carril, palabras in PALABRAS_CLAVE.items():
        for palabra in palabras:
            for pagina in range(MAX_PAGINAS):
                try:
                    avisos = buscar(palabra, pagina)
                except requests.RequestException as e:
                    print(f"  ! '{palabra}' página {pagina}: {e}")
                    break
                finally:
                    time.sleep(PAUSA_SEG)
                recientes = [a for a in avisos if fecha_aviso(a) >= limite]
                for a in recientes:
                    if not pasa_filtros(a):
                        continue
                    previo = encontrados.get(a["id"])
                    if previo:
                        previo["palabras"].add(palabra)
                        previo["carriles"].add(carril)
                    else:
                        encontrados[a["id"]] = {"aviso": a, "palabras": {palabra}, "carriles": {carril}}
                # Resultados vienen del más reciente al más antiguo.
                if len(avisos) < POR_PAGINA or len(recientes) < len(avisos):
                    break
            print(f"  {palabra}: {len(encontrados)} acumuladas")
    return encontrados


def titulo_interesa(titulo):
    t = normalizar(titulo)
    return (not titulo_excluido(titulo)
            and any(contiene_palabra(t, w) for w in portales_empresas.TITULO_INTERES))


def recolectar_portales(dias, postuladas):
    """Filas listas para el Excel con las vacantes de los portales que pasan los filtros."""
    limite = date.today() - timedelta(days=dias)
    filas, claves = [], set()
    for p in portales_empresas.PORTALES:
        n = 0
        try:
            for v in portales_empresas.buscar_portal(p, titulo_interesa, limite):
                if v["fecha_publicacion"] < limite or not ubicacion_valida(v["ubicacion"], v["modalidad"]):
                    continue
                texto = normalizar(v["titulo"] + " " + v["detalle"])
                palabras, carriles = set(), set()
                for carril, lista in PALABRAS_CLAVE.items():
                    for palabra in lista:
                        if contiene_palabra(texto, palabra):
                            palabras.add(palabra)
                            carriles.add(carril)
                if not palabras:
                    if not any(contiene_palabra(normalizar(v["titulo"]), w) for w in TITULO_DATOS):
                        continue
                    carriles.add("Datos-BI")
                clave = clave_vacante(v["titulo"], v["empresa"], v["detalle"])
                if clave in claves:  # algunos portales publican el mismo aviso dos veces
                    continue
                claves.add(clave)
                filas.append(fila_portal(v, palabras, carriles, postuladas))
                n += 1
        except (requests.RequestException, KeyError, ValueError, IndexError) as e:
            print(f"  ! {p['empresa']}: {e}")
        print(f"  Portal {p['empresa']}: {n} pasan los filtros")
    return filas


def empresas_postuladas():
    """Empresas de la hoja CVs de cv_data.xlsx, para marcar vacantes de empresas ya postuladas."""
    ruta = Path(__file__).parent / "cv_data.xlsx"
    if not ruta.exists():
        return set()
    ws = load_workbook(ruta, read_only=True)["CVs"]
    cab = [c for c in next(ws.iter_rows(max_row=1, values_only=True))]
    i = cab.index("empresa")
    return {clave_empresa(r[i]) for r in ws.iter_rows(min_row=2, values_only=True) if r[i]}


def clave_empresa(nombre):
    ruido = {"s", "a", "sa", "sac", "saa", "srl", "eirl", "peru", "del", "de", "la", "el", "cia"}
    return " ".join(p for p in slug(nombre).split("-") if p not in ruido)


def ya_postulada(empresa, postuladas):
    return empresa != "Confidencial" and any(
        p and (p in clave_empresa(empresa) or clave_empresa(empresa) in p) for p in postuladas
    )


def fila_portal(v, palabras, carriles, postuladas):
    return {
        **v,
        "estado": "ya postulaste a esta empresa" if ya_postulada(v["empresa"], postuladas) else "nueva",
        "carril": ", ".join(sorted(carriles)),
        "palabras_clave": ", ".join(sorted(palabras)),
        "encontrada_el": date.today(),
        **analizar_detalle(v["detalle"]),
    }


def fila(datos, postuladas):
    a = datos["aviso"]
    empresa = "Confidencial" if a.get("confidencial") else a.get("empresa")
    return {
        "id": a["id"],
        "estado": "ya postulaste a esta empresa" if ya_postulada(empresa, postuladas) else "nueva",
        "fuente": "Bumeran",
        "fecha_publicacion": fecha_aviso(a),
        "titulo": a.get("titulo"),
        "empresa": empresa,
        "ubicacion": a.get("localizacion"),
        "modalidad": a.get("modalidadTrabajo"),
        "jornada": a.get("tipoTrabajo"),
        "carril": ", ".join(sorted(datos["carriles"])),
        "palabras_clave": ", ".join(sorted(datos["palabras"])),
        "link": f"https://www.bumeran.com.pe/empleos/{slug(a.get('titulo'))}-{a['id']}.html",
        "encontrada_el": date.today(),
        "detalle": a.get("detalle"),
        **analizar_detalle(a.get("detalle")),
    }


COLUMNAS_HISTORIAL = COLUMNAS + ["archivado_el"]


def _preparar_hoja(ws, columnas, color):
    ws.append(columnas)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=color)
    for i, col in enumerate(columnas, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = ANCHOS.get(col, 14)
    ws.freeze_panes = "H2"


def abrir_libro():
    """Devuelve (wb, hoja Vacantes, hoja Historial), creando lo que falte."""
    if ARCHIVO.exists():
        wb = load_workbook(ARCHIVO)
        ws = wb["Vacantes"]
        if [c.value for c in ws[1]] != COLUMNAS:
            sys.exit(f"Las columnas de {ARCHIVO.name} no coinciden con las esperadas; revisa la hoja Vacantes.")
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Vacantes"
        _preparar_hoja(ws, COLUMNAS, "1F3864")
        # postular: "Y" = el usuario postuló; vacío = no
        col = ws.cell(1, COLUMNAS.index("postular") + 1).column_letter
        dv = DataValidation(type="list", formula1='"Y"', allow_blank=True)
        dv.add(f"{col}2:{col}5000")
        ws.add_data_validation(dv)
    if "Historial" in wb.sheetnames:
        hist = wb["Historial"]
        if [c.value for c in hist[1]] != COLUMNAS_HISTORIAL:
            sys.exit(f"Las columnas de la hoja Historial de {ARCHIVO.name} no coinciden con las esperadas.")
    else:
        hist = wb.create_sheet("Historial")
        _preparar_hoja(hist, COLUMNAS_HISTORIAL, "595959")
    return wb, ws, hist


def clave_vacante(titulo, empresa, detalle):
    """Identifica la misma vacante aunque Bumeran la republique con otro id."""
    if not empresa or normalizar(empresa) == "confidencial":
        return "conf|" + slug(titulo) + "|" + slug((detalle or "")[:200])
    # Misma empresa aunque Bumeran y su portal la nombren distinto ("BBVA" / "BBVA Perú").
    empresa = portales_empresas.empresa_canonica(empresa, clave_empresa)
    return clave_empresa(empresa) + "|" + slug(titulo)


def vistos(*hojas):
    """ids y claves de todas las vacantes ya registradas en las hojas dadas."""
    ids, claves = set(), set()
    for h in hojas:
        cab = [c.value for c in h[1]]
        for r in h.iter_rows(min_row=2, values_only=True):
            d = dict(zip(cab, r))
            if d.get("id"):
                ids.add(str(d["id"]))
                claves.add(clave_vacante(d.get("titulo"), d.get("empresa"), d.get("detalle")))
    return ids, claves


def agregar_fila(ws, f, columnas=COLUMNAS):
    ws.append([f.get(c) for c in columnas])
    r = ws.max_row
    for col in ("fecha_publicacion", "encontrada_el", "archivado_el"):
        if col in columnas:
            ws.cell(r, columnas.index(col) + 1).number_format = "yyyy-mm-dd"
    if f.get("link"):
        link = ws.cell(r, columnas.index("link") + 1)
        link.hyperlink = f["link"]
        link.font = Font(color="0563C1", underline="single")
    ws.cell(r, columnas.index("detalle") + 1).alignment = Alignment(wrap_text=False)
    podio = f.get("podio")
    if isinstance(podio, int) and podio <= 20:
        ws.cell(r, 1).fill = PatternFill("solid", fgColor="FFE699" if podio <= 3 else "E2EFDA")
        ws.cell(r, 1).font = Font(bold=True)


def guardar(wb):
    try:
        wb.save(ARCHIVO)
    except PermissionError:
        sys.exit(f"No se pudo guardar {ARCHIVO.name}: ciérralo en Excel y vuelve a correr.")


def filas_de(ws):
    cab = [c.value for c in ws[1]]
    return [dict(zip(cab, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


def pasar_a_portal(ws, f, portal):
    """La fila `f` de Vacantes (venía de Bumeran) pasa a apuntar al portal de la empresa."""
    cab = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, cab.index("id") + 1).value) == str(f["id"]):
            link = ws.cell(r, cab.index("link") + 1)
            link.value = link.hyperlink = portal["link"]
            ws.cell(r, cab.index("fuente") + 1).value = portal["fuente"] + " · también en Bumeran"
            return


def cmd_buscar(dias, portales=True, bumeran=True):
    wb, ws, hist = abrir_libro()
    ids, claves = vistos(ws, hist)
    en_vacantes = {clave_vacante(f.get("titulo"), f.get("empresa"), f.get("detalle")): f for f in filas_de(ws)}
    postuladas = empresas_postuladas()
    nuevas, repetidas, enlazadas, total = [], 0, 0, 0

    # 1. Portales de las empresas (tienen prioridad sobre Bumeran).
    if portales:
        print(f"Buscando en portales de empresas (últimos {dias} días)...")
        for f in recolectar_portales(dias, postuladas):
            total += 1
            clave = clave_vacante(f["titulo"], f["empresa"], f["detalle"])
            previa = en_vacantes.get(clave)
            if previa and previa.get("fuente") == "Bumeran":
                pasar_a_portal(ws, previa, f)
                previa["fuente"] = f["fuente"]
                enlazadas += 1
            if str(f["id"]) in ids or clave in claves:
                repetidas += 1
                continue
            claves.add(clave)
            nuevas.append(f)
    por_clave = {clave_vacante(f["titulo"], f["empresa"], f["detalle"]): f for f in nuevas}

    # 2. Bumeran: si la vacante ya llegó por el portal de la empresa, solo se anota.
    if bumeran:
        print(f"Buscando en Bumeran (últimos {dias} días)...")
        encontrados = recolectar(dias)
        total += len(encontrados)
        for i, d in encontrados.items():
            a = d["aviso"]
            empresa = "Confidencial" if a.get("confidencial") else a.get("empresa")
            clave = clave_vacante(a.get("titulo"), empresa, a.get("detalle"))
            if clave in por_clave:
                por_clave[clave]["fuente"] += " · también en Bumeran"
                repetidas += 1
                continue
            if str(i) in ids or clave in claves:
                repetidas += 1
                continue
            claves.add(clave)
            nuevas.append(fila(d, postuladas))
    nuevas.sort(key=lambda f: f["fecha_publicacion"], reverse=True)

    for f in nuevas:
        agregar_fila(ws, f)
    ws.auto_filter.ref = ws.dimensions
    guardar(wb)
    print(f"\n{total} vacantes pasan los filtros; {repetidas} ya vistas o repetidas entre fuentes; "
          f"{len(nuevas)} nuevas agregadas a {ARCHIVO.name}.")
    if enlazadas:
        print(f"{enlazadas} vacantes que venían de Bumeran ahora apuntan al portal de la empresa.")


def cmd_manual(ruta_json):
    """Agrega una vacante que el usuario pegó en el chat (fuera de Bumeran). Va directo a análisis."""
    import json
    datos = json.loads(Path(ruta_json).read_text(encoding="utf-8"))
    for campo in ("titulo", "empresa", "detalle"):
        if not datos.get(campo):
            sys.exit(f"Falta el campo '{campo}' en {ruta_json}.")
    wb, ws, hist = abrir_libro()
    clave = clave_vacante(datos["titulo"], datos["empresa"], datos["detalle"])
    for nombre, hoja in (("Vacantes", ws), ("Historial", hist)):
        for f in filas_de(hoja):
            if clave_vacante(f.get("titulo"), f.get("empresa"), f.get("detalle")) == clave:
                print(f"YA_VISTA: está en {nombre} (id {f['id']}, estado {f.get('estado')}).")
                sys.exit(2)
    vid = f"M-{date.today():%Y%m%d}-{slug(datos['empresa'])[:24]}"
    ids, _ = vistos(ws, hist)
    n = 2
    while vid in ids:
        vid = f"M-{date.today():%Y%m%d}-{slug(datos['empresa'])[:24]}-{n}"
        n += 1
    fecha = datos.get("fecha_publicacion")
    f = {
        "podio": "Manual",
        "id": vid,
        "estado": "manual: por analizar",
        "fuente": datos.get("fuente") or "Manual (chat)",
        "fecha_publicacion": date.fromisoformat(fecha) if fecha else date.today(),
        "titulo": datos["titulo"],
        "empresa": datos["empresa"],
        "ubicacion": datos.get("ubicacion"),
        "modalidad": datos.get("modalidad"),
        "jornada": datos.get("jornada"),
        "link": datos.get("link"),
        "encontrada_el": date.today(),
        "detalle": datos["detalle"],
        **analizar_detalle(datos["detalle"]),
    }
    agregar_fila(ws, f)
    guardar(wb)
    print(f"AGREGADA: {vid}")


def cmd_archivar(forzar, conservar_con_cv=False, conservar=()):
    """Mueve filas de Vacantes a Historial.

    Sin opciones mueve todo. Con conservar_con_cv deja en Vacantes las filas con cv_id
    (CV ya generado); `conservar` deja además los ids indicados.
    """
    wb, ws, hist = abrir_libro()
    filas = filas_de(ws)
    pendientes = [f for f in filas if str(f.get("postular") or "").strip().upper() == "Y" and not f.get("cv_id")]
    if pendientes and not forzar:
        for f in pendientes:
            print(f"  Y sin CV: {f['id']} · {f['titulo']} · {f['empresa']}")
        sys.exit("Hay vacantes marcadas Y sin cv_id. Genera sus CVs antes de archivar (o usa --forzar).")
    quedan = [f for f in filas if (conservar_con_cv and f.get("cv_id")) or str(f["id"]) in conservar]
    mover = [f for f in filas if f not in quedan]
    for f in mover:
        f["archivado_el"] = date.today()
        agregar_fila(hist, f, COLUMNAS_HISTORIAL)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for f in quedan:
        agregar_fila(ws, f)
    hist.auto_filter.ref = hist.dimensions
    ws.auto_filter.ref = ws.dimensions
    guardar(wb)
    print(f"{len(mover)} vacantes archivadas en Historial; {len(quedan)} quedan en Vacantes.")


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--dias", type=int, default=30, help="antigüedad máxima de la vacante (default 30)")
    fuentes = parser.add_mutually_exclusive_group()
    fuentes.add_argument("--solo-portales", action="store_true", help="no consulta Bumeran")
    fuentes.add_argument("--sin-portales", action="store_true", help="solo Bumeran")
    parser.add_argument("--manual", metavar="JSON", help="agrega una vacante pegada en el chat (titulo, empresa, detalle, ...)")
    parser.add_argument("--archivar", action="store_true", help="mueve todo Vacantes a Historial")
    parser.add_argument("--forzar", action="store_true", help="con --archivar: archiva aunque haya Y sin CV")
    parser.add_argument("--conservar-con-cv", action="store_true", help="con --archivar: deja en Vacantes las filas con cv_id")
    parser.add_argument("--conservar", default="", help="con --archivar: ids separados por coma que se quedan en Vacantes")
    args = parser.parse_args()
    if args.manual:
        cmd_manual(args.manual)
    elif args.archivar:
        conservar = {x.strip() for x in args.conservar.split(",") if x.strip()}
        cmd_archivar(args.forzar, args.conservar_con_cv, conservar)
    else:
        cmd_buscar(args.dias, portales=not args.sin_portales, bumeran=not args.solo_portales)


if __name__ == "__main__":
    main()
