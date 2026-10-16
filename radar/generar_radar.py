"""Genera radar/radar_vacantes.html (artefacto "Radar de vacantes") desde la hoja Vacantes de vacantes.xlsx.

Uso:
    python radar/generar_radar.py

Pestaña "Analizadas": filas con análisis de la skill (recomendacion llena).
Pestaña "Resto": las demás filas, con la calificación rápida y el link.
Las marcas del usuario (Postulé / Analizar) viven en la base de datos del artefacto, no aquí.
"""

import html
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

AQUI = Path(__file__).parent
EXCEL = AQUI.parent / "vacantes.xlsx"
PLANTILLA = AQUI / "plantilla_radar.html"
SALIDA = AQUI / "radar_vacantes.html"


def limpio(texto):
    return html.unescape(re.sub(r"\s+", " ", texto or "")).strip()


def fuentes(texto):
    out = []
    for linea in (texto or "").splitlines():
        if " — " in linea:
            t, u = linea.rsplit(" — ", 1)
            out.append([t.strip(), u.strip()])
    return out


def fecha(v):
    return v.strftime("%d/%m") if hasattr(v, "strftime") else (v or "")


def num(v):
    return v if isinstance(v, (int, float)) else None


def main():
    ws = load_workbook(EXCEL, read_only=True)["Vacantes"]
    filas = list(ws.iter_rows(values_only=True))
    cab = filas[0]
    analizadas, resto = [], []
    for r in filas[1:]:
        f = dict(zip(cab, r))
        if not f.get("id"):
            continue
        base = {
            "id": str(f["id"]), "podio": f.get("podio"), "titulo": f.get("titulo"), "empresa": f.get("empresa"),
            "ubicacion": f.get("ubicacion"), "modalidad": f.get("modalidad"), "fecha": fecha(f.get("fecha_publicacion")),
            "link": f.get("link"), "fuente": f.get("fuente") or "Bumeran", "encaje_rapido": f.get("encaje"),
            "estado": f.get("estado"), "detalle": limpio(f.get("detalle")),
        }
        if f.get("recomendacion"):
            nota_emp = f.get("nota_empresa")
            base.update({
                "recomendacion": f["recomendacion"], "puntaje_global": f.get("puntaje_global"),
                "nota_empresa": num(nota_emp), "nota_puesto": num(f.get("nota_puesto")), "nota_encaje": num(f.get("nota_encaje")),
                "sueldo_esperado": f.get("sueldo_esperado"), "traslado": f.get("traslado"),
                "analisis_empresa": f.get("analisis_empresa"), "analisis_puesto": f.get("analisis_puesto"),
                "analisis_encaje": f.get("analisis_encaje"), "analisis_sueldo": f.get("analisis_sueldo"),
                "veredicto": f.get("veredicto"), "bullets_sugeridos": f.get("bullets_sugeridos"),
                "fuentes": fuentes(f.get("fuentes")),
            })
            analizadas.append(base)
        else:
            base.update({"cubre": f.get("cubre"), "falta": f.get("falta"), "alertas": f.get("alertas")})
            resto.append(base)

    analizadas.sort(key=lambda d: (not isinstance(d["podio"], int), d["podio"] if isinstance(d["podio"], int) else 0,
                                   -(d["puntaje_global"] or 0)))
    orden_grupo = lambda p: {"B": 0, "C": 1, "D": 2}.get(str(p or "")[:1], 3)
    resto.sort(key=lambda d: (orden_grupo(d["podio"]), -(d["encaje_rapido"] or 0)))

    datos = json.dumps({"analizadas": analizadas, "resto": resto}, ensure_ascii=False).replace("</", "<\\/")
    pagina = PLANTILLA.read_text(encoding="utf-8").replace("__DATA__", datos)
    if pagina.count("</script>") != 1:
        sys.exit("La plantilla debe tener exactamente un </script>.")
    SALIDA.write_text(pagina, encoding="utf-8")
    print(f"{SALIDA.name}: {len(analizadas)} analizadas, {len(resto)} en el resto.")


if __name__ == "__main__":
    main()
