#!/usr/bin/env python3
"""Genera CVs personalizados (.docx + .pdf) a partir de cv_data.xlsx.

Uso:
    python generar_cv.py                 # todos los CVs que no estén en estado "enviado"
    python generar_cv.py --cv ID         # solo un CV
    python generar_cv.py --validar       # revisa el Excel sin generar nada
    python generar_cv.py --enviar ID     # copia el CV a enviados/ y lo marca como "enviado"
    python generar_cv.py --forzar        # regenera también los "enviado" (en salida/, nunca en enviados/)
"""
import argparse
import datetime as dt
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips
from openpyxl import load_workbook

# La consola de Windows usa cp1252 y no puede imprimir ✓ ni otros caracteres
sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
EXCEL = BASE / "cv_data.xlsx"
SALIDA = BASE / "salida"
ENVIADOS = BASE / "enviados"

# ---------------------------------------------------------------- diseño
FONT = "Calibri"
AZUL = RGBColor(0x1F, 0x38, 0x64)
AZUL_HEX = "1F3864"
BODY = 9.5                      # pt
PAGE_W, PAGE_H = 11906, 16838   # A4 en twips
MARGIN = 720                    # 0.5"
RIGHT_TAB = PAGE_W - 2 * MARGIN # tab derecha al borde del texto
BULLET_LEFT, BULLET_HANG = 300, 200

ORDEN_SECCIONES_EXTRA = ["Educación", "Proyectos académicos"]  # van después de Experiencia
SECCION_CERT = "Certificaciones"


def tw(n):
    """twips -> Pt (para espaciados)."""
    return Pt(n / 20)


def run(p, text, size=BODY, bold=False, italic=False, color=None):
    r = p.add_run(text)
    f = r.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    if color is not None:
        f.color.rgb = color
    rpr = r._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is not None:
        fonts.set(qn("w:eastAsia"), FONT)
        fonts.set(qn("w:cs"), FONT)
    return r


def para(doc, before=0, after=0, align=None, right_tab=False, border=False):
    p = doc.add_paragraph()
    if border:  # pBdr debe ir antes que tabs/spacing dentro de pPr
        ppr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", AZUL_HEX)):
            b.set(qn(f"w:{k}"), v)
        pbdr.append(b)
        ppr.append(pbdr)
    pf = p.paragraph_format
    if right_tab:
        pf.tab_stops.add_tab_stop(Twips(RIGHT_TAB), WD_TAB_ALIGNMENT.RIGHT)
    pf.space_before = tw(before)
    pf.space_after = tw(after)
    if align is not None:
        p.alignment = align
    return p


def h_seccion(doc, titulo):
    p = para(doc, before=140, after=60, border=True)
    run(p, titulo.upper(), size=10.5, bold=True, color=AZUL)


def h_empresa(doc, empresa, lugar):
    p = para(doc, before=90, after=0, right_tab=True)
    run(p, empresa, size=10, bold=True)
    run(p, "\t" + (lugar or ""))


def h_rol(doc, cargo, fechas):
    p = para(doc, before=30, after=20, right_tab=True)
    run(p, cargo, bold=True, italic=True)
    run(p, "\t" + (fechas or ""), italic=True)


def bullet(doc, texto):
    p = para(doc, after=20)
    pf = p.paragraph_format
    pf.left_indent = Twips(BULLET_LEFT)
    pf.first_line_indent = Twips(-BULLET_HANG)
    pf.tab_stops.add_tab_stop(Twips(BULLET_LEFT))
    run(p, "•\t" + texto)


def nuevo_documento():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(PAGE_W), Twips(PAGE_H)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Twips(MARGIN))
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    # quitar el párrafo vacío inicial si existe
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
    return doc


# ---------------------------------------------------------------- datos
def txt(v):
    return "" if v is None else str(v).strip()


def ids(v):
    return [x.strip() for x in txt(v).replace("\n", ",").replace(";", ",").split(",") if x.strip()]


def hoja(wb, nombre):
    if nombre not in wb.sheetnames:
        sys.exit(f"ERROR: falta la hoja '{nombre}' en {EXCEL.name}")
    filas = list(wb[nombre].iter_rows(values_only=True))
    if not filas:
        return []
    cab = [txt(c) for c in filas[0]]
    out = []
    for f in filas[1:]:
        if all(txt(c) == "" for c in f):
            continue
        out.append({cab[i]: f[i] for i in range(len(cab)) if cab[i]})
    return out


def cargar():
    if not EXCEL.exists():
        sys.exit(f"ERROR: no encuentro {EXCEL}")
    wb = load_workbook(EXCEL, data_only=True)
    perfil = {txt(r["clave"]): txt(r["valor"]) for r in hoja(wb, "Perfil")}
    roles = {txt(r["rol_id"]): r for r in hoja(wb, "Roles")}
    banco = {txt(r["bullet_id"]): r for r in hoja(wb, "Banco")}
    habs = {txt(r["hab_id"]): r for r in hoja(wb, "Habilidades")}
    hab_orden = [txt(r["hab_id"]) for r in hoja(wb, "Habilidades")]
    extras = {txt(r["extra_id"]): r for r in hoja(wb, "Extras")}
    extras_orden = [txt(r["extra_id"]) for r in hoja(wb, "Extras")]
    cvs = hoja(wb, "CVs")
    overrides = {}
    for r in hoja(wb, "Overrides"):
        overrides.setdefault(txt(r["cv_id"]), {})[txt(r["bullet_id"])] = r
    return dict(perfil=perfil, roles=roles, banco=banco, habs=habs, hab_orden=hab_orden,
                extras=extras, extras_orden=extras_orden, cvs=cvs, overrides=overrides)


def resolver(cv, d):
    """Devuelve (estructura del CV, lista de errores)."""
    errores = []
    cid = txt(cv.get("cv_id"))
    ov = d["overrides"].get(cid, {})

    # bullets -> por rol, respetando el orden indicado
    por_rol = {}
    for bid in ids(cv.get("bullets")):
        o = ov.get(bid)
        b = d["banco"].get(bid)
        if o is None and b is None:
            errores.append(f"{cid}: bullet '{bid}' no existe en Banco ni en Overrides")
            continue
        texto = txt(o["texto"]) if o is not None and txt(o.get("texto")) else txt(b["texto"]) if b else ""
        rol = txt(o.get("rol_id")) if o is not None and txt(o.get("rol_id")) else txt(b["rol_id"]) if b else ""
        if not rol:
            errores.append(f"{cid}: bullet '{bid}' no tiene rol_id")
            continue
        if rol not in d["roles"]:
            errores.append(f"{cid}: rol '{rol}' (bullet {bid}) no existe en Roles")
            continue
        if not texto:
            errores.append(f"{cid}: bullet '{bid}' está vacío")
            continue
        por_rol.setdefault(rol, []).append(texto)
    for bid in ov:
        if bid not in ids(cv.get("bullets")):
            errores.append(f"{cid}: override '{bid}' no está listado en la columna bullets (no aparecerá)")

    roles = sorted(por_rol, key=lambda r: float(d["roles"][r].get("orden") or 999))

    hab_ids = ids(cv.get("habilidades")) or d["hab_orden"]
    for h in hab_ids:
        if h not in d["habs"]:
            errores.append(f"{cid}: habilidad '{h}' no existe")
    hab_ids = [h for h in hab_ids if h in d["habs"]]

    ext_ids = ids(cv.get("extras")) or [e for e in d["extras_orden"]
                                         if txt(d["extras"][e].get("por_defecto")).lower() in ("sí", "si", "x", "1", "true")]
    for e in ext_ids:
        if e not in d["extras"]:
            errores.append(f"{cid}: extra '{e}' no existe")
    ext_ids = [e for e in ext_ids if e in d["extras"]]

    if not txt(cv.get("titular")):
        errores.append(f"{cid}: falta titular")
    if not txt(cv.get("resumen")):
        errores.append(f"{cid}: falta resumen")
    return dict(roles=roles, por_rol=por_rol, habs=hab_ids, extras=ext_ids), errores


def construir(cv, est, d):
    P = d["perfil"]
    doc = nuevo_documento()

    p = para(doc, after=20, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, P.get("nombre", "").upper(), size=15, bold=True, color=AZUL)
    p = para(doc, after=30, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, txt(cv["titular"]), size=10.5, bold=True)
    ubic = txt(cv.get("ubicacion")) or P.get("ubicacion", "")
    if ubic == "-":  # "-" = sin ubicación en la cabecera
        ubic = ""
    contacto = " | ".join(x for x in [ubic, P.get("telefono"), P.get("email"), P.get("linkedin"), P.get("portafolio")] if x)
    p = para(doc, after=20, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, contacto)

    h_seccion(doc, "Resumen profesional")
    p = para(doc, after=20, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    run(p, txt(cv["resumen"]))

    h_seccion(doc, "Experiencia")
    empresa_actual = None
    for rid in est["roles"]:
        r = d["roles"][rid]
        if txt(r["empresa"]) != empresa_actual:
            empresa_actual = txt(r["empresa"])
            h_empresa(doc, empresa_actual, txt(r.get("lugar")))
        h_rol(doc, txt(r["cargo"]), txt(r.get("fechas")))
        for b in est["por_rol"][rid]:
            bullet(doc, b)

    extras = [d["extras"][e] for e in est["extras"]]
    for seccion in ORDEN_SECCIONES_EXTRA:
        items = [x for x in extras if txt(x["seccion"]) == seccion]
        if not items:
            continue
        h_seccion(doc, seccion)
        for x in items:
            if seccion == "Educación":
                p = para(doc, before=40, after=10, right_tab=True)
                run(p, txt(x["texto"]), bold=True)
                if txt(x.get("fecha")):
                    run(p, "\t" + txt(x["fecha"]))
                if txt(x.get("detalle")):
                    run(para(doc, after=20), txt(x["detalle"]))
            else:
                bullet(doc, txt(x["texto"]))

    if est["habs"]:
        h_seccion(doc, "Habilidades")
        for h in est["habs"]:
            p = para(doc, after=20)
            run(p, txt(d["habs"][h]["categoria"]) + ": ", bold=True)
            run(p, txt(d["habs"][h]["items"]))

    certs = [x for x in extras if txt(x["seccion"]) == SECCION_CERT]
    if certs:
        h_seccion(doc, "Certificaciones")
        for x in certs:
            bullet(doc, txt(x["texto"]))
    return doc


# ---------------------------------------------------------------- PDF
def buscar_soffice():
    for c in ("soffice", "libreoffice"):
        if shutil.which(c):
            return shutil.which(c)
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(c).exists():
            return c
    return None


def a_pdf(docx_path: Path):
    pdf = docx_path.with_suffix(".pdf")
    try:
        from docx2pdf import convert  # usa Microsoft Word (Windows / Mac)
        convert(str(docx_path), str(pdf))
        if pdf.exists():
            return pdf
    except Exception:
        pass
    so = buscar_soffice()
    if so:
        subprocess.run([so, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        if pdf.exists():
            return pdf
    print("  ! No pude crear el PDF (instala Microsoft Word + docx2pdf, o LibreOffice). Quedó solo el .docx.")
    return None


def paginas(pdf):
    try:
        from pypdf import PdfReader
        return len(PdfReader(str(pdf)).pages)
    except Exception:
        return None


# ---------------------------------------------------------------- comandos
def generar(d, solo=None, forzar=False, validar=False):
    SALIDA.mkdir(exist_ok=True)
    prefijo = d["perfil"].get("prefijo_archivo", "CV")
    filas = [c for c in d["cvs"] if txt(c.get("cv_id"))]
    if solo:
        filas = [c for c in filas if txt(c["cv_id"]) == solo]
        if not filas:
            sys.exit(f"ERROR: no existe el CV '{solo}' en la hoja CVs")
    vistos, total_err = set(), 0
    for cv in filas:
        cid = txt(cv["cv_id"])
        if cid in vistos:
            print(f"✗ {cid}: cv_id duplicado"); total_err += 1; continue
        vistos.add(cid)
        estado = txt(cv.get("estado")).lower()
        est, errores = resolver(cv, d)
        if errores:
            total_err += len(errores)
            for e in errores:
                print("✗", e)
        if validar:
            if not errores:
                print(f"✓ {cid}: OK")
            continue
        if estado == "enviado" and not forzar:
            print(f"- {cid}: enviado, se omite (usa --forzar para regenerarlo en salida/)")
            continue
        if any("no existe" in e or "falta" in e or "vacío" in e or "rol_id" in e for e in errores):
            print(f"  {cid}: no se generó por errores")
            continue
        out = SALIDA / f"{prefijo}_{cid}.docx"
        try:
            construir(cv, est, d).save(out)
        except PermissionError:
            print(f"✗ {cid}: {out.name} está abierto en Word; ciérralo y vuelve a correr"); continue
        pdf = a_pdf(out)
        n = paginas(pdf) if pdf else None
        aviso = f"  ⚠ {n} páginas: hay que recortar" if n and n > 1 else ""
        print(f"✓ {cid}: {out.name}{' + PDF' if pdf else ''}{aviso}")
    return total_err


def enviar(d, cid):
    prefijo = d["perfil"].get("prefijo_archivo", "CV")
    pdf = SALIDA / f"{prefijo}_{cid}.pdf"
    docx = SALIDA / f"{prefijo}_{cid}.docx"
    if not docx.exists():
        sys.exit(f"ERROR: primero genera el CV: python generar_cv.py --cv {cid}")
    hoy = dt.date.today().isoformat()
    dest = ENVIADOS / f"{hoy}_{cid}"
    dest.mkdir(parents=True, exist_ok=True)
    for f in (docx, pdf):
        if f.exists():
            shutil.copy2(f, dest / f.name)
    wb = load_workbook(EXCEL)
    ws = wb["CVs"]
    cab = [txt(c.value) for c in ws[1]]
    ci, ce, cf = cab.index("cv_id") + 1, cab.index("estado") + 1, cab.index("fecha") + 1
    for row in range(2, ws.max_row + 1):
        if txt(ws.cell(row, ci).value) == cid:
            ws.cell(row, ce).value = "enviado"
            ws.cell(row, cf).value = hoy
            break
    try:
        wb.save(EXCEL)
    except PermissionError:
        sys.exit("ERROR: cierra cv_data.xlsx en Excel y vuelve a correr --enviar")
    print(f"✓ {cid} archivado en {dest.relative_to(BASE)} y marcado como enviado")


def main():
    ap = argparse.ArgumentParser(description="Generador de CVs desde cv_data.xlsx")
    ap.add_argument("--cv", help="generar solo este cv_id")
    ap.add_argument("--validar", action="store_true", help="revisar el Excel sin generar")
    ap.add_argument("--enviar", metavar="CV_ID", help="archivar en enviados/ y marcar como enviado")
    ap.add_argument("--forzar", action="store_true", help="regenerar también los enviados (solo en salida/)")
    a = ap.parse_args()
    d = cargar()
    if a.enviar:
        enviar(d, a.enviar)
        return
    errs = generar(d, solo=a.cv, forzar=a.forzar, validar=a.validar)
    if errs:
        print(f"\n{errs} problema(s) encontrados.")
        sys.exit(1)


if __name__ == "__main__":
    main()
