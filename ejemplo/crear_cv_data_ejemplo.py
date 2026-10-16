"""Crea ejemplo/cv_data.xlsx con el perfil ficticio de Camila Rojas Vega.

Uso:
    python ejemplo/crear_cv_data_ejemplo.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SALIDA = Path(__file__).parent / "cv_data.xlsx"

HOJAS = {
    "LEEME": [
        ["Este Excel guarda el texto aprobado de los CVs. Los hechos viven en experiencia_maestra.md."],
        ["Perfil: datos de contacto. Roles: cargos (orden menor = más arriba). Banco: bullets reutilizables."],
        ["Habilidades / Extras: líneas reutilizables. Extras con por_defecto = sí entran si el CV no lista extras."],
        ["CVs: una fila por CV. bullets/habilidades/extras = IDs separados por coma, en el orden deseado."],
        ["Overrides: reescribe un bullet solo para un CV (o crea uno nuevo indicando rol_id). Debe estar listado en CVs.bullets."],
        ["estado: borrador | listo | enviado. Los 'enviado' no se regeneran. Cierra el Excel antes de correr --enviar."],
    ],
    "Perfil": [
        ["clave", "valor"],
        ["nombre", "Camila Rojas Vega"],
        ["ubicacion", "Lima, Perú"],
        ["telefono", "+51 900 000 000"],
        ["email", "camila.rojas@ejemplo.com"],
        ["linkedin", "linkedin.com/in/ejemplo"],
        ["portafolio", ""],
        ["prefijo_archivo", "CV_Camila_Rojas"],
    ],
    "Roles": [
        ["rol_id", "empresa", "lugar", "cargo", "fechas", "orden"],
        ["DA-COM", "Distribuidora Andina S.A.C.", "Lima, Perú", "Analista Comercial", "Mar 2025 – Actualidad", 10],
        ["UNE-PRA", "Universidad Nacional de Ejemplo", "Lima, Perú", "Practicante de Mejora de Procesos", "Ago 2024 – Feb 2025", 20],
    ],
    "Banco": [
        ["bullet_id", "rol_id", "texto", "etiquetas", "notas"],
        ["DA-COM-01", "DA-COM", "Construí en Power BI (Power Query y DAX) el tablero de ventas por canal y vendedor usado por 12 supervisores en la reunión semanal de seguimiento de cuota.", "BI, Power BI, comercial", ""],
        ["DA-COM-02", "DA-COM", "Automaticé con Power Automate el envío diario del reporte de avance de cuota, ahorrando ~1 hora diaria de armado manual en Excel.", "automatización, Power Automate", "La hora es estimación de Camila"],
        ["DA-COM-03", "DA-COM", "Detecté con Python (Pandas) 1,800 clientes duplicados en el maestro del ERP, corrigiendo el conteo de cobertura de la fuerza de ventas.", "Python, calidad de datos", ""],
        ["DA-COM-04", "DA-COM", "Elaboré reportes de ventas por canal y seguimiento de cuotas para la jefatura comercial y la fuerza de ventas.", "comercial, reportes", "Para carril comercial"],
        ["UNE-PRA-01", "UNE-PRA", "Diagnostiqué el proceso de emisión de constancias, que tomaba 10 días frente a un plazo oficial de 3.", "procesos, diagnóstico", ""],
        ["UNE-PRA-02", "UNE-PRA", "Rediseñé el proceso con Microsoft Forms y Power Automate, reduciendo el tiempo de atención de 10 a 4 días.", "automatización, procesos", ""],
    ],
    "Habilidades": [
        ["hab_id", "categoria", "items"],
        ["HAB-BI", "BI y visualización", "Power BI (Power Query, DAX), Excel avanzado (tablas dinámicas, BUSCARX)"],
        ["HAB-DATOS", "Análisis de datos", "Python (Pandas), SQL"],
        ["HAB-AUTO", "Automatización", "Power Automate, Microsoft Forms"],
        ["HAB-COM", "Gestión comercial", "Seguimiento de cuotas, análisis de canales, reportes de ventas"],
    ],
    "Extras": [
        ["extra_id", "seccion", "texto", "fecha", "detalle", "por_defecto"],
        ["EDU-01", "Educación", "Ingeniera Industrial – Universidad Nacional de Ejemplo", "Mar 2019 – Dic 2024", "", "sí"],
        ["CERT-01", "Certificaciones", "Data Analyst in Python – DataCamp, 2025", "", "", "sí"],
    ],
    "CVs": [
        ["cv_id", "estado", "empresa", "puesto", "postulacion_md", "fecha", "ubicacion", "titular", "resumen",
         "bullets", "habilidades", "extras", "notas"],
        ["BASE-DATOS", "listo", "", "Analista de Datos (base)", "", "", "",
         "Analista de Datos | Power BI · Power Automate · Python",
         "Analista de datos con experiencia comercial: construyo tableros en Power BI que usan supervisores de ventas cada semana, automatizo reportes con Power Automate y depuro datos maestros con Python.",
         "DA-COM-01, DA-COM-02, DA-COM-03, UNE-PRA-01, UNE-PRA-02", "HAB-BI, HAB-DATOS, HAB-AUTO", "", "CV base del carril de datos."],
        ["EJEMPLO-BI-RETAIL", "borrador", "Retail Ejemplo S.A.", "Analista de BI", "postulaciones/2026-10-retail-ejemplo-analista-bi.md",
         "", "", "Analista de BI | Power BI · Excel avanzado · Automatización de reportes",
         "Analista de BI con experiencia en el área comercial de una distribuidora: construí el tablero de ventas por canal que usan 12 supervisores y automaticé el reporte diario de avance de cuota. Valido la calidad de los datos maestros con Python.",
         "DA-COM-01, DA-COM-02, DA-COM-03, UNE-PRA-01, UNE-PRA-02", "HAB-BI, HAB-AUTO, HAB-DATOS", "EDU-01, CERT-01",
         "Brechas: BigQuery y supply chain: sin experiencia; no afirmar."],
    ],
    "Overrides": [
        ["cv_id", "bullet_id", "rol_id", "texto"],
        ["EJEMPLO-BI-RETAIL", "DA-COM-01", "DA-COM",
         "Construí en Power BI (Power Query y DAX) el dashboard de KPIs comerciales por canal y vendedor, usado por 12 supervisores en la reunión semanal."],
    ],
}

ANCHOS = {"LEEME": [120], "Perfil": [18, 70], "Roles": [12, 28, 16, 40, 22, 8], "Banco": [14, 12, 90, 30, 40],
          "Habilidades": [14, 30, 80], "Extras": [12, 18, 60, 20, 30, 12],
          "CVs": [20, 11, 18, 30, 40, 12, 14, 45, 70, 45, 30, 20, 50], "Overrides": [20, 14, 12, 100]}


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in HOJAS.items():
        ws = wb.create_sheet(nombre)
        for f in filas:
            ws.append(f)
        if nombre != "LEEME":
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F3864")
        for i, w in enumerate(ANCHOS[nombre], start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    wb.save(SALIDA)
    print(f"Creado {SALIDA}")


if __name__ == "__main__":
    main()
